from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/home/rum/Documents/PlatformIO/Projects/kasif_celebi/imu_doc.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Sayfa ")
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_code(doc, text):
    p = doc.add_paragraph(style="Kod")
    p.paragraph_format.keep_together = True
    p.add_run(text.strip("\n"))
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = widths[index]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cells[col_index].text = str(value)
            cells[col_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[col_index].width = widths[col_index]
            if row_index % 2:
                set_cell_shading(cells[col_index], "EAF2F8")
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
    # Kisa teknik tablolarin satirlarini ayni sayfada tut.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.8)
styles["Normal"].font.color.rgb = RGBColor(25, 25, 25)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12

styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(27)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(16)

for style_name, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(13 if style_name == "Heading 1" else 9)
    style.paragraph_format.space_after = Pt(5)

code_style = styles.add_style("Kod", WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name = "DejaVu Sans Mono"
code_style.font.size = Pt(8.3)
code_style.font.color.rgb = RGBColor(20, 20, 20)
code_style.paragraph_format.left_indent = Inches(0.22)
code_style.paragraph_format.right_indent = Inches(0.22)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(7)
code_style.paragraph_format.line_spacing = 1.0
shd = OxmlElement("w:shd")
shd.set(qn("w:fill"), "F2F2F2")
code_style.element.get_or_add_pPr().append(shd)

for sec in doc.sections:
    header = sec.header.paragraphs[0]
    header.text = "Kâşif Çelebi Robotu   IMU Teknik İnceleme"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(sec.footer.paragraphs[0])

# Kapak
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(55)
p.add_run("KÂŞİF ÇELEBİ ROBOTU").font.size = Pt(11)
p.runs[0].font.bold = True
p.runs[0].font.color.rgb = RGBColor(31, 78, 120)

doc.add_paragraph("IMU Arıza Tespiti ve Doğrudan I2C Çözümü", style="Title")
subtitle = doc.add_paragraph()
subtitle.add_run("MPU6050 uyumsuzluğu için teşhis kaydı ve uygulama kılavuzu").font.size = Pt(15)
subtitle.runs[0].font.color.rgb = RGBColor(70, 70, 70)
subtitle.paragraph_format.space_after = Pt(28)

intro = doc.add_paragraph()
intro.add_run("Ana sonuç  ").bold = True
intro.add_run(
    "IMU kartı I2C hattında 0x68 adresinde cevap verdi, fakat WHO AM I register’ı "
    "beklenen 0x68 yerine 0x72 değerini döndürdü. Adafruit MPU6050 kütüphanesi bu "
    "kimliği reddettiği için imu_ok false kaldı. Sensörü değiştirmek yerine kimlik "
    "kontrolüne bağlı olmayan, doğrudan Wire tabanlı bir sürücü uygulandı. Son testte "
    "imu_ok true oldu ve sensör fiziksel harekete doğru tepki verdi."
)

add_table(
    doc,
    ["Bilgi", "Değer"],
    [
        ["Proje", "Kâşif Çelebi diferansiyel sürüşlü robot"],
        ["Kontrolcü", "Arduino Mega 2560"],
        ["Seri bağlantı", "/dev/ttyUSB0 üzerinden 115200 baud"],
        ["İncelenen sensör", "MPU6050 olarak temin edilen IMU kartı"],
        ["İnceleme tarihi", "11 Eylül 2026"],
        ["Doküman kapsamı", "Teşhis, kanıtlar, kod değişikliği ve doğrulama"],
    ],
    [Inches(1.75), Inches(5.05)],
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.add_run("Dosya adı  ").bold = True
p.add_run("imu_doc.docx")

doc.add_page_break()

doc.add_heading("İçindekiler", level=1)
contents = [
    "1  Sistem ve veri akışı",
    "2  İlk belirti ve olası hata kaynakları",
    "3  ROS ve Arduino ayrıştırması",
    "4  USB seri port kontrolü",
    "5  I2C adres taraması",
    "6  Sensör kimliğinin tespiti",
    "7  Uyumsuz veya klon sensör değerlendirmesi",
    "8  Doğrudan I2C sürücüsünün uygulanması",
    "9  Derleme ve karta yükleme",
    "10  Canlı veri doğrulaması",
    "11  Hareket testi sonuçları",
    "12  Tekrar kullanılabilir teşhis prosedürü",
    "13  Sınırlamalar ve sonraki geliştirmeler",
]
for line in contents:
    p = doc.add_paragraph(line)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(5)

doc.add_heading("Belgenin amacı", level=1)
doc.add_paragraph(
    "Bu belge, USB ve ROS tarafında sıfır görünen IMU verisinin nasıl incelendiğini, "
    "hata alanının nasıl daraltıldığını ve sensörün neden Adafruit MPU6050 kütüphanesiyle "
    "başlatılamadığını adım adım kaydeder. Aynı zamanda uygulanan çözümün register düzeyindeki "
    "ayrıntılarını ve ileride benzer bir arızada izlenecek kontrol sırasını içerir."
)

doc.add_heading("1 Sistem ve veri akışı", level=1)
doc.add_paragraph(
    "Arduino Mega 2560, IMU ve batarya sensörlerini I2C hattından okur. Ölçümleri her 100 ms’de "
    "bir JSON Lines biçiminde USB seri porta gönderir. Jetson üzerinde çalışan Python köprüsü "
    "bu JSON paketlerini ROS 2 mesajlarına dönüştürür. Bu yapı nedeniyle sıfır IMU verisi üç "
    "ayrı katmandan kaynaklanabilir: sensör ve I2C, Arduino firmware’i veya ROS köprüsü."
)
add_table(
    doc,
    ["Katman", "Görev", "Kontrol yöntemi"],
    [
        ["MPU kartı ve I2C", "Ham ivme ve jiroskop ölçümü", "Adres taraması ve register okuma"],
        ["Arduino firmware’i", "Ölçümü JSON paketine dönüştürme", "USB seri porttan cat ile okuma"],
        ["ROS 2 köprüsü", "JSON verisini sensor_msgs Imu mesajına çevirme", "ros2 topic info ve echo"],
    ],
    [Inches(1.65), Inches(2.55), Inches(2.6)],
)

doc.add_heading("2 İlk belirti ve olası hata kaynakları", level=1)
doc.add_paragraph(
    "İlk belirti, araç hareket ettirilmesine rağmen lineer ivme ve açısal hız alanlarının sürekli "
    "sıfır kalmasıydı. İlk aşamada yanlış topic, seri köprü hatası, yanlış birim dönüşümü, açılış "
    "kalibrasyonu, I2C adresi ve donanım bağlantısı olasılıkları birlikte değerlendirildi."
)
add_bullets(doc, [
    "Yanlış ROS topic’inin izlenmesi, gerçek Arduino verisini gizleyebilir.",
    "Aynı seri portun iki süreç tarafından açılması paket kaybına veya erişim hatasına yol açabilir.",
    "mpu.begin çağrısının başarısız olması, firmware’de bütün IMU alanlarının sıfır kalmasına neden olur.",
    "Sensörün 0x68 yerine 0x69 adresinde olması başlatma hatası oluşturabilir.",
    "Kimlik register’ı beklenen değeri döndürmeyen klon veya farklı çipler kütüphane tarafından reddedilebilir.",
])

doc.add_heading("3 ROS ve Arduino ayrıştırması", level=1)
doc.add_paragraph(
    "Projede ROS köprüsü /imu/data_raw topic’ini yayımlar. İlk gözlem /imu/data üzerinde yapıldığı "
    "için topic yayıncısının doğrulanması önerildi. Ancak ham USB verisinde de bütün IMU alanları "
    "sıfır ve imu_ok false görüldüğünde ROS katmanı hata kaynağı olmaktan çıktı."
)
add_code(doc, """
ros2 topic list | grep imu
ros2 topic info /imu/data -v
ros2 topic info /imu/data_raw -v
ros2 topic echo /imu/data_raw --once
""")
doc.add_paragraph(
    "Bu ayrım önemlidir. USB’de doğru, ROS’ta yanlış veri varsa köprü incelenir. USB’de de yanlış "
    "veri varsa önce Arduino ve sensör katmanına dönülür. Bu olayda ikinci durum gerçekleşti."
)

doc.add_heading("4 USB seri port kontrolü", level=1)
doc.add_paragraph(
    "Arduino bağlantısı işletim sistemi, aygıt dosyası ve PlatformIO olmak üzere üç noktadan "
    "doğrulandı. CH340 dönüştürücü 1a86:7523 kimliğiyle görüldü ve /dev/ttyUSB0 aygıtı oluştu."
)
add_code(doc, """
lsusb
ls -l /dev/ttyACM* /dev/ttyUSB* 2>/dev/null
~/.platformio/penv/bin/pio device list
""")
add_table(
    doc,
    ["Kontrol", "Gözlenen sonuç", "Yorum"],
    [
        ["USB aygıtı", "CH340 1a86:7523", "Arduino seri dönüştürücüsü algılandı"],
        ["Seri port", "/dev/ttyUSB0", "Port kullanılabilir durumdaydı"],
        ["Baud hızı", "115200", "Firmware ve terminal ayarı eşleşti"],
        ["Ham çıktı", "JSON Lines", "ROS devre dışı bırakılarak Arduino doğrudan izlendi"],
    ],
    [Inches(1.35), Inches(2.0), Inches(3.45)],
)
add_code(doc, """
stty -F /dev/ttyUSB0 115200 raw -echo
timeout 8s stdbuf -oL cat /dev/ttyUSB0
""")

doc.add_heading("5 I2C adres taraması", level=1)
doc.add_paragraph(
    "Linux terminali Arduino’nun SDA ve SCL hatlarını doğrudan tarayamaz. Bu nedenle firmware’e "
    "1 ile 126 arasındaki bütün I2C adreslerine kısa bir yoklama gönderen tarayıcı eklendi. Bir "
    "aygıt endTransmission çağrısına sıfır hata koduyla cevap verdiğinde adres seri porta yazıldı."
)
add_code(doc, """
void printI2cDevices() {
  Serial.print("{\\\"i2c_addresses\\\":[");
  bool first = true;
  for (uint8_t address = 1; address < 127; address++) {
    Wire.beginTransmission(address);
    if (Wire.endTransmission() == 0) {
      if (!first) Serial.print(',');
      Serial.print(address);
      first = false;
    }
  }
  Serial.println("]}");
}
""")
doc.add_paragraph("Tarama sonucu aşağıdaki gibi alındı:")
add_code(doc, '{"i2c_addresses":[64,104]}')
add_table(
    doc,
    ["Ondalık", "Hexadecimal", "Değerlendirme"],
    [
        ["64", "0x40", "INA219 batarya sensörü"],
        ["104", "0x68", "IMU kartının cevap verdiği adres"],
    ],
    [Inches(1.25), Inches(1.55), Inches(4.0)],
)
doc.add_paragraph(
    "Bu sonuç SDA, SCL ve ortak I2C hattının çalıştığını gösterdi. INA219 zaten doğru ölçüm "
    "üretiyordu. IMU kartının da 0x68 adresinde ACK vermesi, tamamen kopuk kablo veya enerjisiz "
    "kart olasılığını büyük ölçüde dışladı."
)

doc.add_heading("6 Sensör kimliğinin tespiti", level=1)
doc.add_paragraph(
    "I2C adresi, veri yolunda bir aygıtın nerede bulunduğunu gösterir; çip modelini tek başına "
    "kanıtlamaz. Model kontrolü için 0x75 adresindeki WHO AM I register’ı doğrudan okundu."
)
add_code(doc, """
int readI2cRegister(uint8_t address, uint8_t reg) {
  Wire.beginTransmission(address);
  Wire.write(reg);
  if (Wire.endTransmission(false) != 0) return -1;
  if (Wire.requestFrom(address, (uint8_t)1) != 1) return -1;
  return Wire.read();
}

Serial.print("{\\\"imu_who_am_i\\\":");
Serial.print(readI2cRegister(0x68, 0x75));
Serial.println("}");
""")
doc.add_paragraph("Seri port sonucu:")
add_code(doc, '{"imu_who_am_i":114}')
doc.add_paragraph("Ondalık ve hexadecimal dönüşümü terminalde şu komutlarla doğrulandı:")
add_code(doc, """
printf '0x%X\\n' 104
printf '0x%X\\n' 114

# Sonuçlar
# 104 -> 0x68 I2C adresi
# 114 -> 0x72 çip kimliği
""")
add_table(
    doc,
    ["Alan", "Beklenen", "Ölçülen", "Sonuç"],
    [
        ["I2C adresi", "0x68 veya 0x69", "0x68", "Adres geçerli"],
        ["WHO AM I", "MPU6050 için 0x68", "0x72", "Kimlik uyuşmuyor"],
        ["Adafruit begin", "Kimlik eşleşirse true", "false", "Kütüphane sensörü reddetti"],
    ],
    [Inches(1.35), Inches(1.8), Inches(1.3), Inches(2.35)],
)

doc.add_heading("7 Uyumsuz veya klon sensör değerlendirmesi", level=1)
doc.add_paragraph(
    "Kartın üzerinde MPU6050 etiketi bulunması, silikon çipin özgün MPU6050 olduğunu tek başına "
    "kanıtlamaz. Ölçülen 0x72 kimliği MPU6050’nin beklenen 0x68 kimliğiyle uyuşmadığı için kart "
    "uyumsuz, farklı model veya klon olarak değerlendirildi. Bu bulgu sahte ürün şüphesini destekler, "
    "ancak üretici izi veya laboratuvar analizi olmadan kartın ticari anlamda sahte olduğu kesin hükme "
    "bağlanamaz. Teknik olarak kesin olan, çipin Adafruit MPU6050 kimlik kontrolünden geçmediğidir."
)
doc.add_heading("Adafruit kütüphanesinin davranışı", level=2)
doc.add_paragraph(
    "Adafruit MPU6050 kütüphanesi önce I2C aygıtına erişir, ardından WHO AM I değerinin tanımlı "
    "MPU6050 kimliğiyle eşleşmesini ister. Kimlik farklıysa begin çağrısı false döndürür. Proje de "
    "mpuAvailable false kaldığında ölçüm bloğuna girmez ve ax, ay, az, gx, gy, gz alanlarını sıfır gönderir."
)
add_code(doc, """
float ax = 0.0f, ay = 0.0f, az = 0.0f;
float gx = 0.0f, gy = 0.0f, gz = 0.0f;
if (mpuAvailable) {
  // Ölçüm yalnızca başlatma başarılı olduğunda yapılır.
}
""")
doc.add_paragraph(
    "Bu nedenle sıfır değerler sensörün fiziksel olarak hiç cevap vermediğini değil, yazılımın "
    "güvenlik amacıyla ölçüm yolunu kapattığını gösteriyordu."
)

doc.add_heading("8 Doğrudan I2C sürücüsünün uygulanması", level=1)
doc.add_paragraph(
    "Sensörün 0x68 adresinde cevap vermesi ve temel register düzeninin kullanılabilir olması üzerine "
    "Adafruit kimlik kontrolü kaldırıldı. MPU tarafında harici kütüphane yerine Arduino Wire sınıfı "
    "kullanıldı. INA219 kütüphanesi değiştirilmedi."
)
doc.add_heading("Başlatma register’ları", level=2)
add_table(
    doc,
    ["Register", "Yazılan değer", "Amaç"],
    [
        ["0x6B PWR MGT 1", "0x00", "Uyku modundan çıkarma"],
        ["0x19 SMPLRT DIV", "0x09", "Örnekleme bölücüsü"],
        ["0x1A CONFIG", "0x04", "Yaklaşık 20 Hz alçak geçiren filtre"],
        ["0x1B GYRO CONFIG", "0x00", "Jiroskop aralığı artı eksi 250 derece saniye"],
        ["0x1C ACCEL CONFIG", "0x00", "İvmeölçer aralığı artı eksi 2 g"],
    ],
    [Inches(1.7), Inches(1.5), Inches(3.6)],
)
add_code(doc, """
bool initImuDirect(uint8_t address) {
  Wire.beginTransmission(address);
  if (Wire.endTransmission() != 0) return false;

  if (!writeI2cRegister(address, 0x6B, 0x00)) return false;
  delay(100);
  if (!writeI2cRegister(address, 0x19, 0x09)) return false;
  if (!writeI2cRegister(address, 0x1A, 0x04)) return false;
  if (!writeI2cRegister(address, 0x1B, 0x00)) return false;
  if (!writeI2cRegister(address, 0x1C, 0x00)) return false;
  return true;
}
""")

doc.add_heading("Ham ölçümün okunması", level=2)
doc.add_paragraph(
    "Ölçüm bloğu 0x3B adresinden başlar. Tek işlemde 14 bayt okunur: üç ivme ekseni, sıcaklık "
    "ve üç jiroskop ekseni. Her değer iki baytlık işaretli 16 bit sayıdır. Sıcaklık bu projede "
    "kullanılmadığı için iki baytı atlanır."
)
add_code(doc, """
Wire.beginTransmission(mpuAddress);
Wire.write(0x3B);
if (Wire.endTransmission(false) != 0) return false;
if (Wire.requestFrom(mpuAddress, (uint8_t)14) != 14) return false;

int16_t rawAx = ((int16_t)Wire.read() << 8) | Wire.read();
int16_t rawAy = ((int16_t)Wire.read() << 8) | Wire.read();
int16_t rawAz = ((int16_t)Wire.read() << 8) | Wire.read();
""")

doc.add_heading("Birim dönüşümleri", level=2)
doc.add_paragraph(
    "Artı eksi 2 g ayarında ivme ölçeği 16384 LSB g, artı eksi 250 derece saniye ayarında jiroskop "
    "ölçeği 131 LSB derece saniyedir. ROS mesajları SI birimleri istediği için ivme m s kareye, "
    "jiroskop rad s birimine çevrildi."
)
add_code(doc, """
ax = (rawAx / 16384.0f) * 9.80665f;
ay = (rawAy / 16384.0f) * 9.80665f;
az = (rawAz / 16384.0f) * 9.80665f;

const float gyroScale = DEG_TO_RAD / 131.0f;
gx = rawGx * gyroScale;
gy = rawGy * gyroScale;
gz = rawGz * gyroScale;
""")

doc.add_heading("Kalibrasyonun korunması", level=2)
doc.add_paragraph(
    "Mevcut 500 örneklik açılış kalibrasyonu doğrudan okuma fonksiyonuna uyarlandı. Robot bu sırada "
    "düz ve hareketsiz tutulmalıdır. X ve Y ivme ortalamaları ofset olarak çıkarılır. Z ekseninde "
    "yerçekimi korunur; bu nedenle düz konumda az yaklaşık 9,80665 m s kare olmalıdır. Jiroskop "
    "eksenlerinin ortalamaları sıfır ofseti olarak çıkarılır."
)

doc.add_heading("9 Derleme ve karta yükleme", level=1)
doc.add_paragraph(
    "platformio.ini dosyasından Adafruit MPU6050 bağımlılığı kaldırıldı. Adafruit INA219 bağımlılığı "
    "korundu. Proje Arduino Mega 2560 hedefi için derlendi ve /dev/ttyUSB0 üzerinden yüklendi."
)
add_code(doc, """
~/.platformio/penv/bin/pio run

~/.platformio/penv/bin/pio run \\
  --target upload \\
  --upload-port /dev/ttyUSB0
""")
add_table(
    doc,
    ["Doğrulama", "Sonuç"],
    [
        ["Derleme", "Başarılı"],
        ["Flash yazma", "16318 bayt yazıldı"],
        ["Flash doğrulama", "avrdude doğrulaması başarılı"],
        ["RAM kullanımı", "806 bayt ve yaklaşık yüzde 9,8"],
        ["Flash kullanımı", "Yaklaşık yüzde 6,4"],
    ],
    [Inches(2.15), Inches(4.65)],
)

doc.add_heading("10 Canlı veri doğrulaması", level=1)
doc.add_paragraph(
    "Yeni firmware yüklendikten sonra seri port tekrar okundu. Başlangıç paketi sensörün 0x68 "
    "adresinde doğrudan sürücüyle başlatıldığını ve uyumsuz kimlik değeri korunmasına rağmen veri "
    "yolunun açıldığını gösterdi."
)
add_code(doc, """
{"status":"baslatildi","imu_ok":true,
 "imu_address":"0x68","imu_who_am_i":114,"battery_ok":true}
""")
doc.add_paragraph("Düz ve hareketsiz konumdaki örnek ölçüm:")
add_code(doc, """
ax =  0.048226 m/s²
ay =  0.054432 m/s²
az =  9.772343 m/s²
gx = -0.002783 rad/s
gy = -0.000213 rad/s
gz = -0.001251 rad/s
""")
add_table(
    doc,
    ["Alan", "Beklenen düz durum", "Gözlenen aralık", "Değerlendirme"],
    [
        ["ax", "Yaklaşık 0 m/s²", "−0,06 ile +0,10", "Normal gürültü"],
        ["ay", "Yaklaşık 0 m/s²", "−0,09 ile +0,17", "Normal gürültü"],
        ["az", "Yaklaşık +9,806 m/s²", "9,77 ile 9,84", "Yerçekimi doğru"],
        ["gx", "Yaklaşık 0 rad/s", "Yaklaşık ±0,006", "Normal sıfır çevresi"],
        ["gy", "Yaklaşık 0 rad/s", "Yaklaşık ±0,004", "Normal sıfır çevresi"],
        ["gz", "Yaklaşık 0 rad/s", "Yaklaşık ±0,003", "Normal sıfır çevresi"],
    ],
    [Inches(0.8), Inches(1.85), Inches(1.8), Inches(2.35)],
)

doc.add_heading("11 Hareket testi sonuçları", level=1)
doc.add_paragraph(
    "Robot sola ve sağa döndürüldüğünde z ekseni açısal hızı zıt işaretli belirgin tepkiler verdi. "
    "Hareket sonlandığında değer yeniden sıfır çevresine döndü. Bu davranış hem dinamik ölçümün "
    "çalıştığını hem de sensörün hareket yönünü ayırt ettiğini gösterdi."
)
add_table(
    doc,
    ["Durum", "Örnek gz", "Yorum"],
    [
        ["Bir dönüş yönü", "Yaklaşık −1,10 rad/s", "Negatif yön açık biçimde algılandı"],
        ["Ters dönüş yönü", "Yaklaşık +0,99 rad/s", "Pozitif yön açık biçimde algılandı"],
        ["Hareketsiz", "Yaklaşık −0,006 rad/s", "Değer sıfıra geri döndü"],
    ],
    [Inches(2.0), Inches(1.8), Inches(3.0)],
)
doc.add_paragraph(
    "Yatırma testlerinde değerlendirme prensibi farklıdır. Yatırma sırasında ilgili jiroskop "
    "ekseni kısa süreli açısal hız üretir. Araç eğik tutulduğunda yerçekimi ay veya ax eksenine "
    "dağılır ve az büyüklüğü azalır. Düz konuma dönüldüğünde ax ve ay sıfıra, az ise yaklaşık "
    "9,8 m s kareye dönmelidir."
)

doc.add_heading("12 Tekrar kullanılabilir teşhis prosedürü", level=1)
doc.add_paragraph(
    "Aşağıdaki sıra, aynı sistemde ileride görülebilecek sıfır IMU verisini hızlı biçimde ayırmak "
    "için kullanılabilir. Her adım bir önceki katmanı doğrular."
)
add_steps(doc, [
    "Doğru ROS topic’ini ros2 topic list ve ros2 topic info komutlarıyla belirleyin.",
    "ROS köprüsünü durdurun; seri portu tek başına stty ve cat ile okuyun.",
    "JSON paketinde imu_ok alanını kontrol edin. false ise sensör başlatma yoluna geçin.",
    "lsusb ve pio device list ile Arduino portunu doğrulayın.",
    "I2C taramasıyla 0x68 veya 0x69 adresinin cevap verip vermediğini ölçün.",
    "0x75 WHO AM I register’ını okuyun ve değeri beklenen çip kimliğiyle karşılaştırın.",
    "Adres var fakat kimlik uyumsuzsa kütüphane kaynak kodundaki kimlik kontrolünü inceleyin.",
    "Register düzeni uyumluysa doğrudan okuma veya doğru çipe uygun sürücü kullanın.",
    "Düz durumda az yaklaşık 9,8 m s kare ve jiroskop eksenleri yaklaşık sıfır olmalıdır.",
    "İki yönlü dönüş testiyle bir eksende zıt işaretli açısal hızlar görüldüğünü doğrulayın.",
])

doc.add_heading("Hızlı karar tablosu", level=2)
add_table(
    doc,
    ["Gözlem", "En olası alan", "Sonraki işlem"],
    [
        ["Seri port görünmüyor", "USB kablosu, sürücü veya izin", "lsusb, journalctl ve aygıt izinlerini kontrol edin"],
        ["imu_ok false ve I2C adresi yok", "Besleme veya SDA SCL bağlantısı", "Kablolama ve ortak GND kontrolü yapın"],
        ["I2C adresi var ve kimlik yanlış", "Farklı çip veya klon", "Doğru sürücü seçin ya da register uyumluluğunu test edin"],
        ["USB verisi doğru, ROS yanlış", "ROS köprüsü veya topic", "Yayıncıyı ve alan eşlemesini kontrol edin"],
        ["Düz durumda az sıfır", "Okuma veya ölçekleme hatası", "Ham register değerlerini ve ölçeği inceleyin"],
        ["Jiroskop sürekli büyük", "Kalibrasyon veya titreşim", "Hareketsiz açılış ve mekanik sabitleme yapın"],
    ],
    [Inches(2.0), Inches(1.75), Inches(3.05)],
)

doc.add_heading("13 Sınırlamalar ve sonraki geliştirmeler", level=1)
doc.add_paragraph(
    "Doğrudan sürücü ölçüm üretmektedir, ancak 0x72 kimliğinin tam çip modeli kesinleşmemiştir. "
    "Uygulanan ölçekler ve register adresleri MPU6050 uyumluluğuna dayanır. Mevcut sabit ve dinamik "
    "testler bu varsayımın pratikte çalıştığını göstermiştir. Yine de uzun süreli kullanım öncesinde "
    "bilinen açılarla eğim testi ve bilinen açısal hızla ölçek doğrulaması yapılması önerilir."
)
add_bullets(doc, [
    "Açılış kalibrasyonu sırasında robot düz ve tamamen hareketsiz tutulmalıdır.",
    "Sensör daha sonra I2C cevap vermeyi bırakırsa mevcut kod imu_ok değerini false yapar.",
    "Motor gürültüsünü değerlendirmek için motorlar kapalı ve açık durumda ayrı kayıt alınmalıdır.",
    "ROS koordinat düzeni için x ileri, y sol ve z yukarı eksenleri fiziksel montajla doğrulanmalıdır.",
    "Üretim kullanımı için bilinen ve belgelenmiş bir IMU modülü tercih edilmesi bakım riskini azaltır.",
])

doc.add_heading("Sonuç", level=1)
doc.add_paragraph(
    "Hata, ROS mesaj dönüşümünden veya USB seri bağlantısından kaynaklanmıyordu. Arduino sensörü "
    "I2C adresinde görebiliyor, fakat çip 0x72 kimliği döndürdüğü için Adafruit MPU6050 kütüphanesi "
    "başlatmayı reddediyordu. MPU verisi doğrudan Wire üzerinden okununca imu_ok true oldu, düz "
    "konumda yerçekimi doğru ölçüldü ve sağ ile sol dönüşlerde zıt işaretli açısal hızlar elde edildi. "
    "Böylece mevcut JSON ve ROS yapısı korunarak sensör kartının pratikte kullanılabilmesi sağlandı."
)

doc.add_heading("İlgili proje dosyaları", level=2)
add_bullets(doc, [
    "src/main.cpp doğrudan I2C IMU sürücüsü, kalibrasyon ve JSON yayını",
    "platformio.ini Arduino Mega hedefi ve INA219 bağımlılığı",
    "Jetson'daki serial_bridge.py USB JSON verisini ROS 2 mesajlarına dönüştüren köprü",
    "README.md sistem kurulumu, protokol ve sorun giderme komutları",
])

doc.core_properties.title = "IMU Arıza Tespiti ve Doğrudan I2C Çözümü"
doc.core_properties.subject = "Kâşif Çelebi robotu MPU6050 teşhis ve çözüm dokümanı"
doc.core_properties.keywords = "MPU6050, IMU, I2C, Arduino Mega, ROS 2, PlatformIO"
doc.core_properties.comments = "11 Eylül 2026 tarihinde yapılan terminal ve hareket testlerine dayanır."

doc.save(OUT)
print(OUT)
